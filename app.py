import os
import re
import secrets
import hashlib
import uuid
from io import BytesIO
from urllib.parse import quote
from contextlib import asynccontextmanager
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Literal

from fastapi import Depends, FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, ConfigDict, Field, field_validator
from sqlalchemy import func, select
from sqlalchemy.orm import Session
from starlette.middleware.gzip import GZipMiddleware
from starlette.middleware.sessions import SessionMiddleware

from db import (ADMIN_ROLES, ENTRY_ROLES, MEALS, ROLES, AuthLog, ChatMessage, LayoutItem, PriceLog, Setting,
                TableAssignment, UnitPrice, ChatRead, Reservation, ReservationHistory, SessionLocal, TimeSlot,
                User, init_db, now_jst)
from security import hash_password, verify_password

BASE = Path(__file__).parent
IS_PROD = bool(os.environ.get("DATABASE_URL"))
ON_RENDER = bool(os.environ.get("RENDER"))  # Render が自動で設定する
TIME_RE = re.compile(r"^([01]\d|2[0-3]):[0-5]\d$")
INT_MAX = 2_147_483_647  # 人数に上限は設けない(DBの整数型の範囲のみ)


def bootstrap_admin() -> None:
    """ユーザーが1人もいなければ最初の管理者を作る"""
    with SessionLocal() as s:
        if s.scalar(select(func.count()).select_from(User)):
            return
        username = os.environ.get("ADMIN_USERNAME")
        password = os.environ.get("ADMIN_PASSWORD")
        if not (username and password):
            if IS_PROD:
                print("[warn] ユーザー未登録: ADMIN_USERNAME / ADMIN_PASSWORD を設定して再起動してください")
                return
            username, password = "admin", "admin"  # ローカル開発用
            print("[info] ローカル開発用の管理者 admin / admin を作成しました")
        s.add(User(username=username, display_name="管理者", password_hash=hash_password(password),
                   role="admin", is_admin=True))
        s.commit()


def bootstrap_developer() -> None:
    """DEVELOPER_USERNAME / DEVELOPER_PASSWORD が設定され、そのIDが未登録なら最上位ロールのアカウントを作る"""
    username = os.environ.get("DEVELOPER_USERNAME", "").strip()
    password = os.environ.get("DEVELOPER_PASSWORD", "")
    if not (username and password):
        return
    with SessionLocal() as s:
        if s.scalar(select(User).where(User.username == username)):
            return
        s.add(User(username=username, display_name=ROLES["developer"], password_hash=hash_password(password),
                   role="developer", is_admin=True))
        s.commit()
        print(f"[info] 最上位ロールのアカウント {username} を作成しました")


@asynccontextmanager
async def lifespan(_app):
    if ON_RENDER and not IS_PROD:
        # 消えるディスク上の SQLite や開発用 admin/admin で本番起動しないよう止める
        raise RuntimeError("DATABASE_URL が未設定です。Render で PostgreSQL を接続してください")
    init_db()
    bootstrap_admin()
    bootstrap_developer()
    yield


app = FastAPI(title="阿蘇 喫食時間管理表", lifespan=lifespan)
app.add_middleware(
    SessionMiddleware,
    secret_key=os.environ.get("SESSION_SECRET") or secrets.token_hex(32),
    session_cookie="amt_session",
    max_age=60 * 60 * 12,
    same_site="lax",
    https_only=IS_PROD or ON_RENDER,
)
app.mount("/static", StaticFiles(directory=BASE / "static"), name="static")


@app.middleware("http")
async def no_stale_assets(request: Request, call_next):
    """デプロイ後に古い画面(JS/CSS/HTML)が残らないよう、毎回サーバーに更新を確認させる(未変更なら304)"""
    response = await call_next(request)
    if not request.url.path.startswith("/api/"):
        response.headers.setdefault("Cache-Control", "no-cache")
    return response


@app.middleware("http")
async def api_etag(request: Request, call_next):
    """自動更新(ポーリング)の通信量削減: GET の API 応答に ETag を付け、前回と同じ内容なら 304 で本文を送らない"""
    response = await call_next(request)
    if (request.method != "GET" or not request.url.path.startswith("/api/") or response.status_code != 200
            or not response.headers.get("content-type", "").startswith("application/json")):
        return response
    body = b"".join([chunk async for chunk in response.body_iterator])
    etag = f'W/"{hashlib.md5(body).hexdigest()}"'
    not_modified = request.headers.get("if-none-match") == etag
    new = Response(b"" if not_modified else body, status_code=304 if not_modified else 200)
    # Set-Cookie など同名ヘッダーが複数あっても落とさないよう、元のヘッダーをそのまま引き継ぐ
    skip = {b"content-length", b"etag", b"cache-control"}
    new.raw_headers += [(k, v) for k, v in response.raw_headers if k.lower() not in skip]
    new.headers["ETag"] = etag
    new.headers["Cache-Control"] = "private, no-cache"  # ブラウザは毎回確認し、未変更なら手元の結果を使う
    return new


# 応答を gzip 圧縮して送る(JSON・JS・CSS・HTML)
app.add_middleware(GZipMiddleware, minimum_size=500)


# ---------- 共通依存 ----------
def get_db():
    with SessionLocal() as s:
        yield s


def session_user(request: Request, db: Session) -> User | None:
    uid = request.session.get("uid")
    if uid is None:
        return None
    user = db.get(User, uid)
    return user if user and user.active else None


def current_user(request: Request, db: Session = Depends(get_db)) -> User:
    user = session_user(request, db)
    if not user:
        raise HTTPException(401, "ログインしてください")
    return user


def admin_user(user: User = Depends(current_user)) -> User:
    if user.role not in ADMIN_ROLES:
        raise HTTPException(403, "管理者のみ操作できます")
    return user


# ---------- 認証 ----------
class LoginIn(BaseModel):
    username: str
    password: str


@app.post("/api/login")
def login(body: LoginIn, request: Request, db: Session = Depends(get_db)):
    username = body.username.strip()
    user = db.scalar(select(User).where(User.username == username))
    if not user or not user.active or not verify_password(body.password, user.password_hash):
        db.add(AuthLog(user_id=user.id if user else None, username=username[:64], action="login_failed"))
        db.commit()
        raise HTTPException(401, "IDまたはパスワードが違います")
    db.add(AuthLog(user_id=user.id, username=user.username, action="login"))
    db.commit()
    request.session.clear()
    request.session["uid"] = user.id
    return {"ok": True}


@app.get("/logout")
def logout(request: Request, db: Session = Depends(get_db)):
    user = session_user(request, db)
    if user:
        db.add(AuthLog(user_id=user.id, username=user.username, action="logout"))
        db.commit()
    request.session.clear()
    return RedirectResponse("/login", 303)


@app.get("/api/me")
def me(user: User = Depends(current_user)):
    return {"id": user.id, "name": user.display_name, "username": user.username,
            "role": user.role, "role_label": ROLES.get(user.role, ""), "is_admin": user.role in ADMIN_ROLES}


class PasswordIn(BaseModel):
    current: str
    new: str = Field(min_length=4, max_length=128)


@app.post("/api/me/password")
def change_my_password(body: PasswordIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    if not verify_password(body.current, user.password_hash):
        raise HTTPException(400, "現在のパスワードが違います")
    user.password_hash = hash_password(body.new)
    db.add(AuthLog(user_id=user.id, username=user.username, action="password_change"))
    db.commit()
    return {"ok": True}


# ---------- 時間枠 ----------
def slot_labels(db: Session, meal: str) -> list[str]:
    return list(db.scalars(select(TimeSlot.label).where(TimeSlot.meal == meal).order_by(TimeSlot.sort)))


def check_meal(meal: str) -> str:
    if meal not in MEALS:
        raise HTTPException(404, "不明な区分です")
    return meal


@app.get("/api/slots/{meal}")
def get_slots(meal: str, _: User = Depends(current_user), db: Session = Depends(get_db)):
    return slot_labels(db, check_meal(meal))


class SlotsIn(BaseModel):
    slots: list[str] = Field(max_length=48)

    @field_validator("slots")
    @classmethod
    def valid(cls, v):
        v = [s.strip() for s in v if s.strip()]
        for s in v:
            if not TIME_RE.match(s):
                raise ValueError(f"時刻は HH:MM 形式で入力してください: {s}")
        return sorted(set(v))


@app.put("/api/slots/{meal}")
def put_slots(meal: str, body: SlotsIn, _: User = Depends(admin_user), db: Session = Depends(get_db)):
    check_meal(meal)
    for ts in db.scalars(select(TimeSlot).where(TimeSlot.meal == meal)):
        db.delete(ts)
    db.add_all(TimeSlot(meal=meal, label=l, sort=i) for i, l in enumerate(body.slots))
    db.commit()
    return body.slots


# ---------- 台帳 ----------
def norm_time(v: str | None) -> str | None:
    if v in (None, ""):
        return None
    if not TIME_RE.match(v):
        raise ValueError("時刻は HH:MM 形式で入力してください")
    return v


class ReservationIn(BaseModel):
    """編集用(日付・泊数は変更しない)"""
    model_config = ConfigDict(str_strip_whitespace=True)

    room: str = Field(min_length=1, max_length=32)
    guest_name: str = Field(default="", max_length=128)
    adults: int = Field(default=0, ge=0, le=INT_MAX)
    children: int = Field(default=0, ge=0, le=INT_MAX)
    infants: int = Field(default=0, ge=0, le=INT_MAX)
    adult_coupon: int = Field(default=0, ge=0, le=INT_MAX)
    free_adult: int = Field(default=0, ge=0, le=INT_MAX)
    free_child: int = Field(default=0, ge=0, le=INT_MAX)
    child_coupon: int = Field(default=0, ge=0, le=INT_MAX)
    outside: int = Field(default=0, ge=0, le=INT_MAX)
    outside_child: int = Field(default=0, ge=0, le=INT_MAX)
    vip: bool = False
    time_slot: str | None = None
    allergy: str = Field(default="", max_length=2000)
    note: str = Field(default="", max_length=2000)
    grouped: bool = False           # グループ登録する
    group_with: int | None = None   # 紐づける相手の予約ID

    _time = field_validator("time_slot")(norm_time)


class ReservationCreateIn(ReservationIn):
    """登録用: date は管理表で表示中の日付。泊数分の日付に1件ずつ登録する"""
    date: date
    nights: int = Field(default=1, ge=1, le=30)


class TimeSlotIn(BaseModel):
    time_slot: str | None = None

    _time = field_validator("time_slot")(norm_time)


TRACKED = ("date", "nights", "night_no", "time_slot", "room", "guest_name", "adults", "children", "infants",
           "adult_coupon", "child_coupon", "free_adult", "free_child", "outside", "outside_child", "vip", "allergy", "note", "group_id", "entered_at")
GROUP_FIELDS = {"grouped", "group_with"}


def iso(v: datetime | None) -> str | None:
    return v.isoformat() if v else None


def snapshot(r: Reservation) -> dict:
    return {f: (v.isoformat() if isinstance(v, (date, datetime)) else v)
            for f in TRACKED for v in [getattr(r, f)]}


def table_names(db: Session, rids) -> dict[int, list[str]]:
    """予約ID → アサインされたテーブル番号(番号順)"""
    rids = list(rids)
    out: dict[int, list[str]] = {}
    if not rids:
        return out
    for rid, name in db.execute(select(TableAssignment.reservation_id, LayoutItem.name)
                                .join(LayoutItem, LayoutItem.id == TableAssignment.table_id)
                                .where(TableAssignment.reservation_id.in_(rids))):
        out.setdefault(rid, []).append(name)
    for v in out.values():
        v.sort(key=natural_key)
    return out


def natural_key(s: str):
    """T2 < T10 のように数字部分を数として並べる"""
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", s)]


def res_out(db: Session, r: Reservation) -> dict:
    return to_dict(r, user_names(db), table_names(db, [r.id]).get(r.id, []))


def to_dict(r: Reservation, names: dict[int, str], tables: list[str] | None = None) -> dict:
    return {
        "id": r.id, "meal": r.meal, **snapshot(r), "tables": tables or [],
        "created_at": iso(r.created_at), "created_by": names.get(r.created_by, ""),
        "updated_at": iso(r.updated_at), "updated_by": names.get(r.updated_by, ""),
        "deleted": r.deleted_at is not None,
        "entered_by": names.get(r.entered_by, ""),
        "deleted_at": iso(r.deleted_at), "deleted_by": names.get(r.deleted_by, ""),
    }


def user_names(db: Session) -> dict[int, str]:
    return dict(db.execute(select(User.id, User.display_name)).all())


def get_reservation(db: Session, meal: str, rid: int, *, editable: bool = True) -> Reservation:
    r = db.get(Reservation, rid)
    if not r or r.meal != meal:
        raise HTTPException(404, "見つかりません")
    if editable and r.deleted_at is not None:
        raise HTTPException(400, "削除済みの予約は変更できません")
    return r


def record(db: Session, r: Reservation, user: User, action: str, changes: dict) -> None:
    """更新日時・更新者を記録し、履歴を1件追記する"""
    now = now_jst()
    r.updated_at, r.updated_by = now, user.id
    db.flush()
    db.add(ReservationHistory(reservation_id=r.id, action=action, changes=changes, changed_at=now,
                              changed_by=user.id))


def change(db: Session, r: Reservation, values: dict, user: User) -> None:
    """値を更新し、差分があれば履歴に残す(commit は呼び出し側)"""
    before = snapshot(r)
    for k, v in values.items():
        setattr(r, k, v)
    after = snapshot(r)
    diff = {f: [before[f], after[f]] for f in TRACKED if before[f] != after[f]}
    if diff:
        record(db, r, user, "update", diff)


# ---------- グループ ----------
def group_target(db: Session, meal: str, d: date, target_id: int | None, self_id: int | None) -> Reservation:
    t = db.get(Reservation, target_id) if target_id else None
    if not t or t.meal != meal or t.date != d or t.deleted_at is not None or t.id == self_id:
        raise HTTPException(400, "紐づける予約を選んでください")
    return t


def ensure_group(db: Session, t: Reservation, user: User) -> str:
    """相手がまだグループでなければ新しいグループにする"""
    if not t.group_id:
        change(db, t, {"group_id": uuid.uuid4().hex}, user)
    return t.group_id


def shrink_group(db: Session, group_id: str | None, user: User) -> None:
    """抜けた結果1件だけ残ったグループは解消する"""
    if not group_id:
        return
    db.flush()
    rest = list(db.scalars(select(Reservation).where(Reservation.group_id == group_id,
                                                     Reservation.deleted_at.is_(None))))
    if len(rest) == 1:
        change(db, rest[0], {"group_id": None}, user)


@app.get("/api/{meal}/reservations")
def list_reservations(meal: str, d: date, include_deleted: bool = False, _: User = Depends(current_user),
                      db: Session = Depends(get_db)):
    check_meal(meal)
    q = select(Reservation).where(Reservation.meal == meal, Reservation.date == d)
    if not include_deleted:
        q = q.where(Reservation.deleted_at.is_(None))
    names = user_names(db)
    rows = list(db.scalars(q.order_by(Reservation.room, Reservation.id)))
    tables = table_names(db, (r.id for r in rows))
    return [to_dict(r, names, tables.get(r.id)) for r in rows]


@app.post("/api/{meal}/reservations")
def create_reservation(meal: str, body: ReservationCreateIn, user: User = Depends(current_user),
                       db: Session = Depends(get_db)):
    check_meal(meal)
    now = now_jst()
    values = body.model_dump(exclude={"date", "nights"} | GROUP_FIELDS)
    stay_id = uuid.uuid4().hex if body.nights > 1 else None
    target = group_target(db, meal, body.date, body.group_with, None) if body.grouped else None
    created = []
    for i in range(body.nights):
        d = body.date + timedelta(days=i)
        # 連泊時は、相手の同じ日の予約(相手も連泊なら)と紐づける
        mate = target if i == 0 or not target else (
            db.scalar(select(Reservation).where(Reservation.stay_id == target.stay_id, Reservation.date == d,
                                                Reservation.deleted_at.is_(None))) if target.stay_id else None)
        group_id = ensure_group(db, mate, user) if mate else None
        r = Reservation(meal=meal, date=d, nights=body.nights, night_no=i + 1,
                        stay_id=stay_id, group_id=group_id, **values,
                        created_at=now, created_by=user.id, updated_at=now, updated_by=user.id)
        db.add(r)
        db.flush()
        record(db, r, user, "create", {f: [None, v] for f, v in snapshot(r).items()})
        created.append(r)
    db.commit()
    names = user_names(db)
    return [to_dict(r, names) for r in created]


@app.put("/api/{meal}/reservations/{rid}")
def update_reservation(meal: str, rid: int, body: ReservationIn, user: User = Depends(current_user),
                       db: Session = Depends(get_db)):
    r = get_reservation(db, check_meal(meal), rid)
    values = body.model_dump(exclude=GROUP_FIELDS)
    old_group = r.group_id
    if not body.grouped:
        values["group_id"] = None
    else:
        t = group_target(db, meal, r.date, body.group_with, r.id)
        if not (r.group_id and t.group_id == r.group_id):
            values["group_id"] = ensure_group(db, t, user)
    change(db, r, values, user)
    if old_group and old_group != r.group_id:
        shrink_group(db, old_group, user)
    db.commit()
    return res_out(db, r)


@app.patch("/api/{meal}/reservations/{rid}/time")
def set_time(meal: str, rid: int, body: TimeSlotIn, user: User = Depends(current_user),
             db: Session = Depends(get_db)):
    r = get_reservation(db, check_meal(meal), rid)
    change(db, r, {"time_slot": body.time_slot}, user)
    db.commit()
    return res_out(db, r)


class EnteredIn(BaseModel):
    entered: bool


@app.patch("/api/{meal}/reservations/{rid}/entered")
def set_entered(meal: str, rid: int, body: EnteredIn, user: User = Depends(current_user),
                db: Session = Depends(get_db)):
    """ステータス(入場済)の切り替え。レストランと最上位ロールのみ"""
    if user.role not in ENTRY_ROLES:
        raise HTTPException(403, "入場済の操作はレストランのみ可能です")
    r = get_reservation(db, check_meal(meal), rid)
    if body.entered != (r.entered_at is not None):
        change(db, r, {"entered_at": now_jst() if body.entered else None,
                       "entered_by": user.id if body.entered else None}, user)
        db.commit()
    return res_out(db, r)


@app.delete("/api/{meal}/reservations/{rid}")
def delete_reservation(meal: str, rid: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    """論理削除: データは残し、削除日時・削除者を記録する"""
    r = get_reservation(db, check_meal(meal), rid)
    r.deleted_at, r.deleted_by = now_jst(), user.id
    record(db, r, user, "delete", {})
    db.commit()
    return res_out(db, r)


@app.post("/api/{meal}/reservations/{rid}/restore")
def restore_reservation(meal: str, rid: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    r = get_reservation(db, check_meal(meal), rid, editable=False)
    if r.deleted_at is None:
        raise HTTPException(400, "削除されていない予約です")
    r.deleted_at, r.deleted_by = None, None
    record(db, r, user, "restore", {})
    db.commit()
    return res_out(db, r)


@app.get("/api/{meal}/reservations/{rid}/history")
def reservation_history(meal: str, rid: int, _: User = Depends(current_user), db: Session = Depends(get_db)):
    get_reservation(db, check_meal(meal), rid, editable=False)
    names = user_names(db)
    rows = db.scalars(select(ReservationHistory).where(ReservationHistory.reservation_id == rid)
                      .order_by(ReservationHistory.id.desc()))
    return [{"action": h.action, "changes": h.changes, "changed_at": iso(h.changed_at),
             "changed_by": names.get(h.changed_by, "")} for h in rows]


# ---------- ユーザー管理 ----------
def user_dict(u: User) -> dict:
    return {"id": u.id, "username": u.username, "display_name": u.display_name,
            "role": u.role, "active": u.active}


Role = Literal["developer", "admin", "front", "restaurant"]


def set_role(u: User, role: str) -> None:
    u.role, u.is_admin = role, role in ADMIN_ROLES


def check_role_change(me_: User, target_role: str | None, target: User | None = None) -> None:
    """最上位ロールのアカウントを作る・編集する・最上位ロールにするのは最上位ロールのみ"""
    if me_.role == "developer":
        return
    if target_role == "developer" or (target and target.role == "developer"):
        raise HTTPException(403, f"{ROLES['developer']} のアカウントは {ROLES['developer']} のみ操作できます")


class UserCreateIn(BaseModel):
    username: str = Field(pattern=r"^[A-Za-z0-9_.-]{2,64}$")
    display_name: str = Field(min_length=1, max_length=64)
    password: str = Field(min_length=4, max_length=128)
    role: Role = "front"


class UserUpdateIn(BaseModel):
    display_name: str | None = Field(default=None, min_length=1, max_length=64)
    password: str | None = Field(default=None, min_length=4, max_length=128)
    role: Role | None = None
    active: bool | None = None


@app.get("/api/users")
def list_users(_: User = Depends(admin_user), db: Session = Depends(get_db)):
    return [user_dict(u) for u in db.scalars(select(User).order_by(User.id))]


@app.post("/api/users")
def create_user(body: UserCreateIn, me_: User = Depends(admin_user), db: Session = Depends(get_db)):
    check_role_change(me_, body.role)
    if db.scalar(select(User).where(User.username == body.username)):
        raise HTTPException(400, "そのログインIDは既に使われています")
    u = User(username=body.username, display_name=body.display_name.strip(),
             password_hash=hash_password(body.password))
    set_role(u, body.role)
    db.add(u)
    db.commit()
    return user_dict(u)


@app.put("/api/users/{uid}")
def update_user(uid: int, body: UserUpdateIn, me_: User = Depends(admin_user), db: Session = Depends(get_db)):
    u = db.get(User, uid)
    if not u:
        raise HTTPException(404, "見つかりません")
    check_role_change(me_, body.role, u)
    if u.id == me_.id and ((body.role and body.role != me_.role) or body.active is False):
        raise HTTPException(400, "自分自身のロール変更・無効化はできません")
    if body.display_name is not None:
        u.display_name = body.display_name.strip()
    if body.password is not None:
        u.password_hash = hash_password(body.password)
    if body.role is not None:
        set_role(u, body.role)
    if body.active is not None:
        u.active = body.active
    db.commit()
    return user_dict(u)


# ---------- チャット ----------
CHAT_LIMIT = 200


class ChatIn(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True)

    body: str = Field(default="", max_length=2000)
    reservation_id: int | None = None


class ReadIn(BaseModel):
    last_id: int


def chat_dict(m: ChatMessage, users: dict[int, User], res: dict[int, Reservation]) -> dict:
    u = users.get(m.user_id)
    r = res.get(m.reservation_id)
    retracted = m.retracted_at is not None
    return {
        "id": m.id, "user_id": m.user_id, "name": u.display_name if u else "",
        "role": m.role, "role_label": ROLES.get(m.role, ""),
        "body": "" if retracted else m.body, "retracted": retracted,
        "created_at": iso(m.created_at),
        "reservation": None if retracted or not r else {
            "id": r.id, "meal": r.meal, "date": r.date.isoformat(), "room": r.room,
            "guest_name": r.guest_name, "time_slot": r.time_slot, "deleted": r.deleted_at is not None},
    }


def unread_count(db: Session, user: User) -> int:
    last = db.scalar(select(ChatRead.last_read_id).where(ChatRead.user_id == user.id)) or 0
    return db.scalar(select(func.count()).select_from(ChatMessage).where(
        ChatMessage.id > last, ChatMessage.user_id != user.id, ChatMessage.retracted_at.is_(None))) or 0


@app.get("/api/chat/messages")
def chat_messages(user: User = Depends(current_user), db: Session = Depends(get_db)):
    msgs = list(db.scalars(select(ChatMessage).order_by(ChatMessage.id.desc()).limit(CHAT_LIMIT)))[::-1]
    users = {u.id: u for u in db.scalars(select(User))}
    rids = {m.reservation_id for m in msgs if m.reservation_id}
    res = {r.id: r for r in db.scalars(select(Reservation).where(Reservation.id.in_(rids)))} if rids else {}
    reads = [{"user_id": uid, "name": users[uid].display_name, "role_label": ROLES.get(users[uid].role, ""),
              "last_read_id": last}
             for uid, last in db.execute(select(ChatRead.user_id, ChatRead.last_read_id))
             if uid in users and users[uid].active]
    return {"me": user.id, "messages": [chat_dict(m, users, res) for m in msgs], "reads": reads}


@app.post("/api/chat/messages")
def post_chat(body: ChatIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    if not body.body and not body.reservation_id:
        raise HTTPException(400, "メッセージを入力してください")
    r = db.get(Reservation, body.reservation_id) if body.reservation_id else None
    if body.reservation_id and not r:
        raise HTTPException(400, "添付した予約が見つかりません")
    m = ChatMessage(user_id=user.id, role=user.role, body=body.body, reservation_id=body.reservation_id)
    db.add(m)
    db.flush()
    mark_read(db, user, m.id)
    db.commit()
    return {"id": m.id}


@app.post("/api/chat/messages/{mid}/retract")
def retract_chat(mid: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    m = db.get(ChatMessage, mid)
    if not m or m.user_id != user.id:
        raise HTTPException(404, "取り消せるメッセージが見つかりません")
    if m.retracted_at is None:
        m.retracted_at = now_jst()
        db.commit()
    return {"ok": True}


def mark_read(db: Session, user: User, last_id: int) -> None:
    cr = db.get(ChatRead, user.id)
    if not cr:
        db.add(ChatRead(user_id=user.id, last_read_id=last_id))
    elif last_id > cr.last_read_id:
        cr.last_read_id = last_id


@app.post("/api/chat/read")
def chat_read(body: ReadIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    latest = db.scalar(select(func.max(ChatMessage.id))) or 0
    mark_read(db, user, min(body.last_id, latest))
    db.commit()
    return {"unread": unread_count(db, user)}


@app.get("/api/chat/unread")
def chat_unread(user: User = Depends(current_user), db: Session = Depends(get_db)):
    return {"unread": unread_count(db, user)}


# ---------- 集計 ----------
MEAL_LABELS = {"dinner": "夕食", "breakfast": "朝食"}
WEEKDAYS = "月火水木金土日"
# (キー, 見出し) — 画面と Excel の列順
SUMMARY_COLS = [
    ("groups", "組数"), ("adults", "大人"), ("children", "子供"), ("infants", "幼児"), ("total", "計"),
    ("adult_coupon", "大人クーポン(食事付)"), ("child_coupon", "子供クーポン(食事付)"),
    ("free_adult", "フリー大人(生打ち)"), ("free_child", "フリー子供(生打ち)"),
    ("outside", "外来大人"), ("outside_child", "外来子供"),
    ("entered", "入場済(組)"), ("vip", "VIP(組)"),
]
VIP_GUESTS_LABEL = "VIPのお客様"  # 数値列の後ろに置く文字列の列
SUMMED = ("adults", "children", "infants", "adult_coupon", "child_coupon", "free_adult", "free_child", "outside",
          "outside_child")


def summarize(db: Session, meal: str, start: date, end: date) -> dict:
    if end < start:
        raise HTTPException(400, "終了日は開始日以降にしてください")
    if (end - start).days > 366:
        raise HTTPException(400, "期間は1年以内にしてください")
    days = {start + timedelta(days=i): dict.fromkeys((k for k, _ in SUMMARY_COLS), 0)
            for i in range((end - start).days + 1)}
    rows = db.scalars(select(Reservation).where(
        Reservation.meal == meal, Reservation.date >= start, Reservation.date <= end,
        Reservation.deleted_at.is_(None)))  # 削除済みは集計しない
    vips = {d: [] for d in days}
    for r in rows:
        a = days[r.date]
        if r.vip:
            a["vip"] += 1
            vips[r.date].append(r)
        a["groups"] += 1
        for k in SUMMED:
            a[k] += getattr(r, k)
        a["total"] += r.adults + r.children + r.infants
        a["entered"] += r.entered_at is not None
    # VIPのお客様: 時間順(未定は最後)、同じ時間なら部屋番号の数値順
    room_key = lambda room: [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", room)]
    vip_names = {d: "、".join(f"★{r.room} {r.guest_name}".strip() for r in
                             sorted(v, key=lambda r: (r.time_slot or "99:99", room_key(r.room))))
                 for d, v in vips.items()}
    total = {k: sum(a[k] for a in days.values()) for k, _ in SUMMARY_COLS}
    return {
        "meal": meal, "start": start.isoformat(), "end": end.isoformat(),
        "columns": [{"key": k, "label": l} for k, l in SUMMARY_COLS],
        "vip_label": VIP_GUESTS_LABEL,
        "days": [{"date": d.isoformat(), "weekday": WEEKDAYS[d.weekday()], **a, "vip_guests": vip_names[d]}
                 for d, a in days.items()],
        "total": total,
    }


@app.get("/api/{meal}/summary")
def meal_summary(meal: str, start: date, end: date, _: User = Depends(current_user), db: Session = Depends(get_db)):
    return summarize(db, check_meal(meal), start, end)


@app.get("/api/{meal}/summary.xlsx")
def meal_summary_xlsx(meal: str, start: date, end: date, _: User = Depends(current_user),
                      db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    data = summarize(db, check_meal(meal), start, end)
    title = f"{MEAL_LABELS[meal]}集計"
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append([f"{title}　{start:%Y/%m/%d}〜{end:%Y/%m/%d}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"出力日時 {now_jst():%Y/%m/%d %H:%M}(削除済みの予約は除く)"])
    ws["A2"].font = Font(size=9, color="777777")
    ws.append([])
    head = ["日付", "曜日"] + [l for _, l in SUMMARY_COLS] + [VIP_GUESTS_LABEL]
    ws.append(head)
    thin = Side(style="thin", color="BBBBBB")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E6F1FB")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    # 合計行は見出しの直後(5行目)
    ws.append(["合計", ""] + [data["total"][k] for k, _ in SUMMARY_COLS] + [""])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="F2F2F2")
        c.border = Border(top=thin, left=thin, right=thin, bottom=Side(style="medium", color="8FA9C4"))
    for d in data["days"]:
        ws.append([date.fromisoformat(d["date"]), d["weekday"]] + [d[k] for k, _ in SUMMARY_COLS] + [d["vip_guests"]])
        row = ws[ws.max_row]
        row[0].number_format = "yyyy/mm/dd"
        color = {"土": "185FA5", "日": "C0392B"}.get(d["weekday"])
        for c in row:
            c.border = border
            if color and c.column <= 2:
                c.font = Font(color=color)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 5
    for i in range(3, len(head) + 1):
        ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = 11
    vip_col = ws.cell(row=4, column=len(head)).column_letter
    ws.column_dimensions[vip_col].width = 40
    for c in ws[vip_col][4:]:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[4].height = 32
    ws.freeze_panes = "C6"  # 見出しと合計行を固定
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:5"  # 印刷時は各ページに見出しと合計行

    buf = BytesIO()
    wb.save(buf)
    fname = f"{title}_{start:%Y%m%d}-{end:%Y%m%d}.xlsx"
    return Response(buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})


# ---------- 設定値(店舗マスタ) ----------
SETTING_DEFAULTS = {"dinner_stay_minutes": "120"}  # 夕食の滞在時間(分): テーブルの重なりの判定に使う


def get_setting(db: Session, key: str) -> str:
    row = db.get(Setting, key)
    return row.value if row else SETTING_DEFAULTS.get(key, "")


def stay_minutes(db: Session, meal: str) -> int:
    try:
        return max(15, int(get_setting(db, f"{meal}_stay_minutes") or 120))
    except ValueError:
        return 120


@app.get("/api/settings")
def get_settings(_: User = Depends(current_user), db: Session = Depends(get_db)):
    return {"dinner_stay_minutes": stay_minutes(db, "dinner")}


class SettingsIn(BaseModel):
    dinner_stay_minutes: int = Field(ge=15, le=600)


@app.put("/api/settings")
def put_settings(body: SettingsIn, user: User = Depends(admin_user), db: Session = Depends(get_db)):
    old = stay_minutes(db, "dinner")
    row = db.get(Setting, "dinner_stay_minutes") or Setting(key="dinner_stay_minutes")
    row.value = str(body.dinner_stay_minutes)
    db.merge(row)
    if old != body.dinner_stay_minutes:  # 操作ログ(店舗マスタ)に残す
        db.add(PriceLog(user_id=user.id, meal="dinner", item="stay_minutes", old_price=old,
                        new_price=body.dinner_stay_minutes, effective_from=now_jst().date()))
    db.commit()
    return get_settings(user, db)


# ---------- レイアウト(見取り図) ----------
LAYOUT_W, LAYOUT_H = 3000, 2000  # 見取り図の大きさ(画面の単位)


def layout_dict(it: LayoutItem) -> dict:
    return {"id": it.id, "kind": it.kind, "shape": it.shape, "name": it.name, "seats": it.seats,
            "x": it.x, "y": it.y, "w": it.w, "h": it.h}


@app.get("/api/layout/{meal}")
def get_layout(meal: str, _: User = Depends(current_user), db: Session = Depends(get_db)):
    check_meal(meal)
    items = db.scalars(select(LayoutItem).where(LayoutItem.meal == meal, LayoutItem.deleted_at.is_(None))
                       .order_by(LayoutItem.id))
    return {"width": LAYOUT_W, "height": LAYOUT_H, "items": [layout_dict(i) for i in items],
            "default_sizes": default_sizes(db, meal)}


# テーブルを追加するときの大きさ(形ごと)。店舗の設定として保存する
DEFAULT_SIZES = {"round": [80, 80], "square": [80, 80], "rect": [160, 80]}


def default_sizes(db: Session, meal: str) -> dict:
    import json
    sizes = {k: list(v) for k, v in DEFAULT_SIZES.items()}
    try:
        saved = json.loads(get_setting(db, f"{meal}_table_sizes") or "{}")
        for k, v in saved.items():
            if k in sizes and isinstance(v, list) and len(v) == 2:
                sizes[k] = [int(v[0]), int(v[1])]
    except (ValueError, TypeError):
        pass
    return sizes


class SizeIn(BaseModel):
    w: int = Field(ge=20, le=1000, multiple_of=20)
    h: int = Field(ge=20, le=1000, multiple_of=20)


class SizesIn(BaseModel):
    round: SizeIn
    square: SizeIn
    rect: SizeIn


@app.put("/api/layout/{meal}/default-sizes")
def put_default_sizes(meal: str, body: SizesIn, _: User = Depends(admin_user), db: Session = Depends(get_db)):
    import json
    check_meal(meal)
    body.round.h = body.round.w  # 丸は幅=高さ(直径)
    key = f"{meal}_table_sizes"
    row = db.get(Setting, key) or Setting(key=key)
    row.value = json.dumps({k: [getattr(body, k).w, getattr(body, k).h] for k in DEFAULT_SIZES})
    db.merge(row)
    db.commit()
    return default_sizes(db, meal)


class LayoutItemIn(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True)

    id: int | None = None
    kind: Literal["table", "marker"]
    shape: Literal["round", "square", "rect", "label"] = "square"
    name: str = Field(default="", max_length=32)
    seats: int = Field(default=0, ge=0, le=999)
    x: int = Field(ge=0, le=LAYOUT_W)
    y: int = Field(ge=0, le=LAYOUT_H)
    w: int = Field(ge=10, le=LAYOUT_W)
    h: int = Field(ge=10, le=LAYOUT_H)


class LayoutIn(BaseModel):
    items: list[LayoutItemIn] = Field(max_length=2000)


@app.get("/api/layout/{meal}/usage")
def layout_usage(meal: str, _: User = Depends(admin_user), db: Session = Depends(get_db)):
    """テーブルごとの今日以降のアサイン件数(レイアウト編集で削除するときの確認用)"""
    check_meal(meal)
    today = now_jst().date()
    rows = db.execute(select(TableAssignment.table_id, func.count())
                      .join(Reservation, Reservation.id == TableAssignment.reservation_id)
                      .where(Reservation.date >= today, Reservation.deleted_at.is_(None))
                      .group_by(TableAssignment.table_id))
    return {str(tid): n for tid, n in rows}


@app.put("/api/layout/{meal}")
def put_layout(meal: str, body: LayoutIn, user: User = Depends(admin_user), db: Session = Depends(get_db)):
    """見取り図をまるごと保存。無くなったテーブルは論理削除し、今日以降のアサインを解除する"""
    check_meal(meal)
    names = [i.name for i in body.items if i.kind == "table"]
    if any(not n for n in names):
        raise HTTPException(400, "テーブル番号が空のテーブルがあります")
    dup = sorted({n for n in names if names.count(n) > 1}, key=natural_key)
    if dup:
        raise HTTPException(400, f"テーブル番号が重複しています: {'、'.join(dup)}")
    current = {i.id: i for i in db.scalars(select(LayoutItem).where(LayoutItem.meal == meal,
                                                                   LayoutItem.deleted_at.is_(None)))}
    keep = set()
    for it in body.items:
        values = it.model_dump(exclude={"id"})
        if it.kind == "marker":
            values.update(shape="label", seats=0)
        row = current.get(it.id) if it.id else None
        if row:
            for k, v in values.items():
                setattr(row, k, v)
            keep.add(row.id)
        else:
            db.add(LayoutItem(meal=meal, **values))
    today = now_jst().date()
    removed = 0
    for rid_, row in current.items():
        if rid_ in keep:
            continue
        row.deleted_at = now_jst()
        # 今日以降のアサインは解除(過去のものは記録として残す)
        for a, r in db.execute(select(TableAssignment, Reservation)
                               .join(Reservation, Reservation.id == TableAssignment.reservation_id)
                               .where(TableAssignment.table_id == row.id, Reservation.date >= today)).all():
            before = ", ".join(table_names(db, [r.id]).get(r.id, []))
            db.delete(a)
            db.flush()
            after = ", ".join(table_names(db, [r.id]).get(r.id, []))
            record(db, r, user, "update", {"tables": [before, after]})
            removed += 1
    db.commit()
    return {**get_layout(meal, user, db), "unassigned": removed}


# ---------- テーブルアサイン ----------
def slot_minutes(t: str | None) -> int | None:
    if not t:
        return None
    h, m = t.split(":")
    return int(h) * 60 + int(m)


@app.get("/api/{meal}/assign")
def assign_state(meal: str, d: date, _: User = Depends(current_user), db: Session = Depends(get_db)):
    """その日のアサインの状態: 予約(削除済みを除く)・アサイン・滞在時間"""
    check_meal(meal)
    rows = list(db.scalars(select(Reservation).where(Reservation.meal == meal, Reservation.date == d,
                                                     Reservation.deleted_at.is_(None))))
    ids = [r.id for r in rows]
    assigns = list(db.scalars(select(TableAssignment).where(TableAssignment.reservation_id.in_(ids)))) if ids else []
    names = user_names(db)
    tables = table_names(db, ids)
    return {
        "date": d.isoformat(), "stay_minutes": stay_minutes(db, meal), "slots": slot_labels(db, meal),
        "reservations": [to_dict(r, names, tables.get(r.id)) for r in rows],
        "assignments": [{"reservation_id": a.reservation_id, "table_id": a.table_id} for a in assigns],
    }


class AssignIn(BaseModel):
    table_id: int


def table_conflict(db: Session, r: Reservation, table_id: int) -> Reservation | None:
    """同じ日に、そのテーブルを滞在時間が重なる別の予約が使っていれば返す"""
    stay = stay_minutes(db, r.meal)
    start = slot_minutes(r.time_slot)
    for other in db.scalars(select(Reservation).join(TableAssignment, TableAssignment.reservation_id == Reservation.id)
                            .where(TableAssignment.table_id == table_id, Reservation.date == r.date,
                                   Reservation.meal == r.meal, Reservation.id != r.id,
                                   Reservation.deleted_at.is_(None))):
        o = slot_minutes(other.time_slot)
        if o is not None and start < o + stay and o < start + stay:
            return other
    return None


@app.post("/api/{meal}/reservations/{rid}/tables")
def assign_table(meal: str, rid: int, body: AssignIn, user: User = Depends(current_user),
                 db: Session = Depends(get_db)):
    r = get_reservation(db, check_meal(meal), rid)
    if not r.time_slot:
        raise HTTPException(400, "時間が未定の予約はアサインできません。先に時間を決めてください")
    t = db.get(LayoutItem, body.table_id)
    if not t or t.meal != meal or t.kind != "table" or t.deleted_at is not None:
        raise HTTPException(404, "テーブルが見つかりません")
    if db.scalar(select(TableAssignment).where(TableAssignment.reservation_id == rid,
                                               TableAssignment.table_id == t.id)):
        return res_out(db, r)
    other = table_conflict(db, r, t.id)
    if other:
        raise HTTPException(409, f"{t.name} は {other.time_slot} の {other.room} {other.guest_name} 様が使用中です"
                                 f"(滞在時間 {stay_minutes(db, meal)}分)")
    before = ", ".join(table_names(db, [r.id]).get(r.id, []))
    db.add(TableAssignment(reservation_id=r.id, table_id=t.id, created_by=user.id))
    db.flush()
    after = ", ".join(table_names(db, [r.id]).get(r.id, []))
    record(db, r, user, "update", {"tables": [before, after]})
    db.commit()
    return res_out(db, r)


@app.delete("/api/{meal}/reservations/{rid}/tables/{table_id}")
def unassign_table(meal: str, rid: int, table_id: int, user: User = Depends(current_user),
                   db: Session = Depends(get_db)):
    r = get_reservation(db, check_meal(meal), rid)
    a = db.scalar(select(TableAssignment).where(TableAssignment.reservation_id == rid,
                                                TableAssignment.table_id == table_id))
    if a:
        before = ", ".join(table_names(db, [r.id]).get(r.id, []))
        db.delete(a)
        db.flush()
        after = ", ".join(table_names(db, [r.id]).get(r.id, []))
        record(db, r, user, "update", {"tables": [before, after]})
        db.commit()
    return res_out(db, r)


# ---------- 店舗マスタ(単価) ----------
# 単価を設定する項目(画面・帳票の並び順)
PRICE_ITEMS = [
    ("adults", "大人"), ("children", "子供"), ("infants", "幼児"),
    ("adult_coupon", "大人クーポン(食事付)"), ("child_coupon", "子供クーポン(食事付)"),
    ("free_adult", "フリー大人(生打ち)"), ("free_child", "フリー子供(生打ち)"),
    ("outside", "外来大人"), ("outside_child", "外来子供"),
]
PRICE_KEYS = [k for k, _ in PRICE_ITEMS]


def prices_on(db: Session, meal: str, d: date) -> dict[str, int]:
    """その日に有効な単価(未設定は0円)"""
    rows = db.scalars(select(UnitPrice).where(UnitPrice.meal == meal, UnitPrice.effective_from <= d)
                      .order_by(UnitPrice.effective_from, UnitPrice.id))
    prices = dict.fromkeys(PRICE_KEYS, 0)
    for p in rows:  # 適用開始日の古い順に上書き → 最新が残る
        if p.item in prices:
            prices[p.item] = p.price
    return prices


@app.get("/api/prices")
def get_prices(_: User = Depends(admin_user), db: Session = Depends(get_db)):
    today = now_jst().date()
    names = user_names(db)
    labels = dict(PRICE_ITEMS)
    out = {"items": [{"key": k, "label": l} for k, l in PRICE_ITEMS], "today": today.isoformat(), "meals": {}}
    for meal in MEALS:
        hist = db.scalars(select(UnitPrice).where(UnitPrice.meal == meal)
                          .order_by(UnitPrice.effective_from.desc(), UnitPrice.id.desc()))
        upcoming = {}
        for p in db.scalars(select(UnitPrice).where(UnitPrice.meal == meal, UnitPrice.effective_from > today)
                            .order_by(UnitPrice.effective_from)):
            upcoming.setdefault(p.item, {"price": p.price, "from": p.effective_from.isoformat()})
        out["meals"][meal] = {
            "current": prices_on(db, meal, today),
            "upcoming": upcoming,  # 今日より後に適用される予定の単価(項目ごとに直近のもの)
            "history": [{"id": p.id, "item": p.item, "label": labels.get(p.item, p.item), "price": p.price,
                         "from": p.effective_from.isoformat(), "at": iso(p.created_at),
                         "by": names.get(p.created_by, "")} for p in hist],
        }
    return out


class PricesIn(BaseModel):
    meal: str
    effective_from: date
    prices: dict[str, int]

    @field_validator("prices")
    @classmethod
    def valid(cls, v):
        for k, p in v.items():
            if k not in PRICE_KEYS:
                raise ValueError(f"不明な項目です: {k}")
            if not 0 <= p <= INT_MAX:
                raise ValueError("単価は0円以上で入力してください")
        return v


@app.put("/api/prices")
def put_prices(body: PricesIn, user: User = Depends(admin_user), db: Session = Depends(get_db)):
    """適用開始日つきで単価を登録。その日に有効な単価から変わる項目だけ記録する(同じ適用開始日なら上書き)"""
    check_meal(body.meal)
    before = prices_on(db, body.meal, body.effective_from)
    changed = 0
    for item, price in body.prices.items():
        same_day = db.scalar(select(UnitPrice).where(UnitPrice.meal == body.meal, UnitPrice.item == item,
                                                     UnitPrice.effective_from == body.effective_from))
        old = same_day.price if same_day else before[item]
        if old == price:
            continue
        if same_day:
            same_day.price, same_day.created_by, same_day.created_at = price, user.id, now_jst()
        else:
            db.add(UnitPrice(meal=body.meal, item=item, price=price, effective_from=body.effective_from,
                             created_by=user.id))
        db.add(PriceLog(user_id=user.id, meal=body.meal, item=item, old_price=old, new_price=price,
                        effective_from=body.effective_from))
        changed += 1
    db.commit()
    return {"changed": changed}


# ---------- 帳票(1日分の詳細) ----------
def report_data(db: Session, meal: str, d: date) -> dict:
    prices = prices_on(db, meal, d)
    rows = list(db.scalars(select(Reservation).where(
        Reservation.meal == meal, Reservation.date == d, Reservation.deleted_at.is_(None))))  # 削除済みは除く
    room_key = lambda room: [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", room)]
    rows.sort(key=lambda r: (r.time_slot or "99:99", room_key(r.room)))
    counts = {k: sum(getattr(r, k) for r in rows) for k in PRICE_KEYS}
    items = [{"key": k, "label": l, "count": counts[k], "price": prices[k], "amount": counts[k] * prices[k]}
             for k, l in PRICE_ITEMS]
    tables = table_names(db, (r.id for r in rows))
    details = [{
        "id": r.id, "time_slot": r.time_slot, "room": r.room, "guest_name": r.guest_name, "vip": r.vip,
        "tables": tables.get(r.id, []),
        "nights": r.nights, "night_no": r.night_no, "entered": r.entered_at is not None,
        **{k: getattr(r, k) for k in PRICE_KEYS},
        "amount": sum(getattr(r, k) * prices[k] for k in PRICE_KEYS),
        "allergy": r.allergy, "note": r.note,
    } for r in rows]
    return {
        "meal": meal, "date": d.isoformat(), "weekday": WEEKDAYS[d.weekday()],
        "items": items, "rows": details,
        "total": {"amount": sum(i["amount"] for i in items), "groups": len(rows),
                  **counts},
    }


@app.get("/api/{meal}/report")
def meal_report(meal: str, d: date, _: User = Depends(current_user), db: Session = Depends(get_db)):
    return report_data(db, check_meal(meal), d)


@app.get("/api/{meal}/report.xlsx")
def meal_report_xlsx(meal: str, d: date, _: User = Depends(current_user), db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    data = report_data(db, check_meal(meal), d)
    title = f"{MEAL_LABELS[meal]}帳票"
    wb = Workbook()
    ws = wb.active
    ws.title = title
    thin = Side(style="thin", color="BBBBBB")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    head_fill = PatternFill("solid", fgColor="E6F1FB")
    total_fill = PatternFill("solid", fgColor="F2F2F2")
    yen = '#,##0'

    def add_row(values, fill=None, bold=False):
        """1行追加し、値を入れた列だけに罫線・塗り・太字を付ける"""
        ws.append(values)
        for c in ws[ws.max_row][:len(values)]:
            c.border = border
            if fill:
                c.fill = fill
            if bold:
                c.font = Font(bold=True)

    ws.append([f"{title}　{d:%Y/%m/%d}({data['weekday']})"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"出力日時 {now_jst():%Y/%m/%d %H:%M}(削除済みの予約は除く。単価はこの日に有効なもの)"])
    ws["A2"].font = Font(size=9, color="777777")
    ws.append([])
    ws.append(["■ 項目別集計"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    add_row(["項目", "人数", "単価", "金額"], head_fill, True)
    for it in data["items"]:
        add_row([it["label"], it["count"], it["price"], it["amount"]])
        ws.cell(row=ws.max_row, column=3).number_format = yen
        ws.cell(row=ws.max_row, column=4).number_format = yen
    add_row(["合計", "", "", data["total"]["amount"]], total_fill, True)  # 人数は大人と内訳で重なるため合計しない
    ws.cell(row=ws.max_row, column=4).number_format = yen

    ws.append([])
    ws.append(["■ 予約明細"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    labels = [l for _, l in PRICE_ITEMS]
    has_tables = meal == "dinner"  # テーブルアサインは夕食だけ
    add_row(["時間", "部屋", "代表者名", "VIP", *labels, "金額", "アレルギー", "備考"] + (["テーブル"] if has_tables else []),
            head_fill, True)
    head_row = ws.max_row
    for c in ws[head_row]:
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in data["rows"]:
        add_row([r["time_slot"] or "未定", r["room"], r["guest_name"], "★" if r["vip"] else "",
                 *[r[k] for k in PRICE_KEYS], r["amount"], r["allergy"], r["note"]]
                + ([", ".join(r["tables"])] if has_tables else []))
        ws.cell(row=ws.max_row, column=4 + len(PRICE_KEYS) + 1).number_format = yen
    add_row(["合計", "", f"{data['total']['groups']}組", "", *[data["total"][k] for k in PRICE_KEYS],
             data["total"]["amount"], "", ""], total_fill, True)
    ws.cell(row=ws.max_row, column=4 + len(PRICE_KEYS) + 1).number_format = yen

    widths = {"A": 22, "B": 9, "C": 18, "D": 5}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for i in range(5, 5 + len(PRICE_KEYS)):
        ws.column_dimensions[ws.cell(row=head_row, column=i).column_letter].width = 10
    amount_col = ws.cell(row=head_row, column=5 + len(PRICE_KEYS)).column_letter
    ws.column_dimensions[amount_col].width = 12
    ws.row_dimensions[head_row].height = 32
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buf = BytesIO()
    wb.save(buf)
    fname = f"{title}_{d:%Y%m%d}.xlsx"
    return Response(buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})


# ---------- 操作ログ ----------
FIELD_LABELS = {
    "date": "日付", "nights": "泊数", "night_no": "何泊目", "time_slot": "時間", "room": "部屋",
    "guest_name": "代表者名", "adults": "大人", "children": "子供", "infants": "幼児",
    "adult_coupon": "大人クーポン(食事付)", "free_adult": "フリー大人(生打ち)",
    "child_coupon": "子供クーポン(食事付)", "free_child": "フリー子供(生打ち)", "outside": "外来大人", "outside_child": "外来子供", "vip": "VIP",
    "allergy": "アレルギー", "note": "備考", "group_id": "グループ", "entered_at": "ステータス",
    "tables": "テーブル",
}
RES_ACTIONS = {"create": "登録", "update": "変更", "delete": "削除", "restore": "復元"}
AUTH_ACTIONS = {"login": "ログイン", "login_failed": "ログイン失敗", "logout": "ログアウト",
                "password_change": "パスワード変更"}
LOG_MAX_DAYS = 93


def fmt_value(field: str, v) -> str:
    if field == "time_slot":
        return v or "未定"
    if field in ("group_id", "vip"):
        return "あり" if v else "なし"
    if field == "entered_at":
        return "入場済" if v else "空白"
    if field == "tables":
        return v or "なし"
    return "(空欄)" if v in (None, "") else str(v)


def change_detail(action: str, changes: dict) -> str:
    if action == "create":  # 登録時は入力された項目だけ
        return "、".join(f"{FIELD_LABELS.get(f, f)} {fmt_value(f, b)}" for f, (_, b) in changes.items()
                        if b not in (None, "", 0) and f not in ("date", "nights", "night_no", "room", "guest_name"))
    return "、".join(f"{FIELD_LABELS.get(f, f)}: {fmt_value(f, a)} → {fmt_value(f, b)}" for f, (a, b) in changes.items())


@app.get("/api/logs")
def operation_logs(start: date, end: date, _: User = Depends(current_user), db: Session = Depends(get_db)):
    """予約・チャット・ログインの操作を新しい順にまとめて返す"""
    if end < start:
        raise HTTPException(400, "終了日は開始日以降にしてください")
    if (end - start).days >= LOG_MAX_DAYS:
        raise HTTPException(400, f"期間は{LOG_MAX_DAYS}日以内にしてください")
    t0 = datetime.combine(start, datetime.min.time())
    t1 = datetime.combine(end + timedelta(days=1), datetime.min.time())
    users = {u.id: u for u in db.scalars(select(User))}

    def who(uid, fallback=""):
        u = users.get(uid)
        return {"user": u.display_name if u else fallback, "role": u.role if u else "",
                "role_label": ROLES.get(u.role, "") if u else ""}

    logs = []
    # 予約
    for h, r in db.execute(select(ReservationHistory, Reservation)
                           .join(Reservation, Reservation.id == ReservationHistory.reservation_id)
                           .where(ReservationHistory.changed_at >= t0, ReservationHistory.changed_at < t1)):
        logs.append({
            "_id": h.id, "at": iso(h.changed_at), "kind": "reservation", "kind_label": "予約",
            "action": "入場済" if set(h.changes) == {"entered_at"} else RES_ACTIONS.get(h.action, h.action),
            "target": f"{MEAL_LABELS.get(r.meal, '')} {r.date:%m/%d} {r.room} {r.guest_name}".strip(),
            "link": f"/{r.meal}?d={r.date.isoformat()}&hl={r.id}",
            "detail": change_detail(h.action, h.changes or {}), **who(h.changed_by)})
    # チャット(取り消されたメッセージの本文はログにも出さない)
    chat_q = select(ChatMessage).where(
        ((ChatMessage.created_at >= t0) & (ChatMessage.created_at < t1)) |
        ((ChatMessage.retracted_at >= t0) & (ChatMessage.retracted_at < t1)))
    for m in db.scalars(chat_q):
        base = {"_id": m.id, "kind": "chat", "kind_label": "チャット", "target": "チャット", "link": "/chat", **who(m.user_id)}
        body = "(取り消し済み)" if m.retracted_at else m.body
        if t0 <= m.created_at < t1:
            logs.append({**base, "at": iso(m.created_at), "action": "投稿", "detail": body})
        if m.retracted_at and t0 <= m.retracted_at < t1:
            logs.append({**base, "_id": m.id + 0.5, "at": iso(m.retracted_at), "action": "取り消し", "detail": ""})
    # ログイン
    for a in db.scalars(select(AuthLog).where(AuthLog.at >= t0, AuthLog.at < t1)):
        logs.append({"_id": a.id, "at": iso(a.at), "kind": "auth", "kind_label": "ログイン",
                     "action": AUTH_ACTIONS.get(a.action, a.action), "target": "", "link": "",
                     "detail": f"ログインID: {a.username}" if a.action == "login_failed" else "",
                     "failed": a.action == "login_failed", **who(a.user_id, a.username)})
    # 店舗マスタ(単価変更)
    price_labels = dict(PRICE_ITEMS)
    yen = lambda n: "―" if n is None else f"{n:,}円"
    for p in db.scalars(select(PriceLog).where(PriceLog.at >= t0, PriceLog.at < t1)):
        if p.item == "stay_minutes":
            logs.append({"_id": p.id, "at": iso(p.at), "kind": "master", "kind_label": "店舗マスタ", "action": "設定変更",
                         "target": f"{MEAL_LABELS.get(p.meal, '')} 滞在時間", "link": "/store-master",
                         "detail": f"{p.old_price}分 → {p.new_price}分", **who(p.user_id)})
            continue
        logs.append({"_id": p.id, "at": iso(p.at), "kind": "master", "kind_label": "店舗マスタ", "action": "単価変更",
                     "target": f"{MEAL_LABELS.get(p.meal, '')} {price_labels.get(p.item, p.item)}", "link": "/store-master",
                     "detail": f"{yen(p.old_price)} → {yen(p.new_price)}({p.effective_from:%Y/%m/%d}〜)",
                     **who(p.user_id)})
    # 同じ秒の操作は記録順(各テーブルのID順)で並べる
    logs.sort(key=lambda x: (x["at"], x.pop("_id")), reverse=True)
    return logs


# ---------- ページ ----------
PAGES = {"dinner", "breakfast", "dinner-summary", "breakfast-summary", "dinner-report", "breakfast-report",
         "table-assign", "layout-edit", "chat", "logs", "admin", "store-master"}
ADMIN_PAGES = {"admin", "store-master", "layout-edit"}  # 管理者・developer だけが開けるページ


def file_response(request: Request, path: Path, **kwargs) -> Response:
    """ファイルを返す。ブラウザが持っているものと同じなら中身を送らず 304 だけ返す(通信量対策)"""
    st = path.stat()
    etag = f'"{st.st_mtime_ns:x}-{st.st_size:x}"'
    headers = {"ETag": etag, "Cache-Control": "no-cache", **kwargs.pop("headers", {})}
    if etag in request.headers.get("if-none-match", ""):
        return Response(status_code=304, headers=headers)
    return FileResponse(path, headers=headers, **kwargs)


@app.get("/favicon.ico", include_in_schema=False)
def favicon(request: Request):
    return file_response(request, BASE / "static" / "favicon.ico", media_type="image/x-icon")


@app.get("/healthz")
def healthz():
    return {"ok": True}


@app.get("/")
def root():
    return RedirectResponse("/dinner")


@app.get("/login")
def login_page(request: Request, db: Session = Depends(get_db)):
    if session_user(request, db):
        return RedirectResponse("/dinner")
    return file_response(request, BASE / "pages" / "login.html")


@app.get("/{page}")
def page(page: str, request: Request, db: Session = Depends(get_db)):
    if page not in PAGES:
        return RedirectResponse("/dinner")
    user = session_user(request, db)
    if not user:
        return RedirectResponse("/login")
    if page in ADMIN_PAGES and user.role not in ADMIN_ROLES:
        return RedirectResponse("/dinner")
    return file_response(request, BASE / "pages" / f"{page.replace('-', '_')}.html")
